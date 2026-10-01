"""Plain Word reports built from persisted evidence; no external template or service."""

import hashlib
import io
import json
import re
from datetime import datetime, timezone

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


def clean(value):
    # Source telemetry is untrusted and may contain XML-forbidden control bytes.
    return re.sub(r"[^\x09\x0a\x0d\x20-\ud7ff\ue000-\ufffd\U00010000-\U0010ffff]", "", str(value if value is not None else "—"))


def _document(title):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = section.bottom_margin = Inches(0.65)
    section.left_margin = section.right_margin = Inches(0.7)
    for name in ("Normal", "Title", "Heading 1", "Heading 2"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.color.rgb = RGBColor(0, 0, 0)
        for border in style.element.xpath(".//w:pBdr"):
            border.getparent().remove(border)
    doc.styles["Normal"].font.size = Pt(10)
    doc.styles["Normal"].paragraph_format.space_after = Pt(6)
    doc.core_properties.author = "IdentityTrace"
    doc.core_properties.title = title
    doc.add_paragraph(title, style="Title")
    doc.add_paragraph("Generated " + datetime.now(timezone.utc).strftime("%Y-%m-%d %H:%M UTC"))
    return doc


def _table(doc, headings, rows, widths):
    table = doc.add_table(rows=1, cols=len(headings))
    table.autofit = False
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        node = OxmlElement("w:" + edge)
        for key, value in (("val", "single"), ("sz", "4"), ("color", "D9D9D9")):
            node.set(qn("w:" + key), value)
        borders.append(node)
    table._tbl.tblPr.append(borders)
    for col, width in zip(table.columns, widths, strict=True):
        col.width = Inches(width)
    for row_index, values in enumerate([headings, *rows]):
        row = table.rows[0] if row_index == 0 else table.add_row()
        for cell, value, width in zip(row.cells, values, widths, strict=True):
            cell.width = Inches(width)
            cell.text = clean(value)
            if row_index == 0:
                cell.paragraphs[0].runs[0].bold = True
                shade = OxmlElement("w:shd")
                shade.set(qn("w:fill"), "E7E7E7")
                cell._tc.get_or_add_tcPr().append(shade)
    repeat = OxmlElement("w:tblHeader")
    table.rows[0]._tr.get_or_add_trPr().append(repeat)
    doc.add_paragraph()


def _bytes(doc):
    output = io.BytesIO()
    doc.save(output)
    return output.getvalue()


def summary_report(incidents, total, filters):
    doc = _document("IdentityTrace Incident Review")
    doc.add_paragraph(f"This report contains {len(incidents)} of {total} matching incidents. "
                      "Review the evidence and record a disposition before treating a finding as confirmed compromise.")
    doc.add_paragraph("Filters: " + clean(json.dumps(filters, default=str, sort_keys=True)))
    _table(doc, ["Incident and identity", "Severity", "Status", "Last event UTC"], [
        [f"{i['correlation_rule_title']}\n{i['identity_id']}\n{i['incident_id']}",
         i["severity"], i["status"] + (" (superseded)" if i["superseded_at"] else ""), i["last_event_at"]]
        for i in incidents], [3.1, 0.8, 1.1, 1.6])
    if not incidents:
        doc.add_paragraph("No incidents matched the selected filters.")
    return _bytes(doc)


def incident_report(incident, events):
    doc = _document("IdentityTrace Incident Report")
    doc.add_paragraph(clean(incident["correlation_rule_title"]))
    doc.add_paragraph("This report records the finding, analyst disposition and evidence at export time. "
                      "The score and confidence are heuristics, not probabilities of compromise.")
    _table(doc, ["Field", "Value"], [
        ["Incident", incident["incident_id"]], ["Identity", incident["identity_id"]],
        ["Severity and status", f"{incident['severity']} / {incident['status']}"],
        ["Classification", incident["classification"]],
        ["Score and confidence", f"{incident['score']} / {incident['confidence']}"],
        ["First and last event", f"{incident['first_event_at']} to {incident['last_event_at']}"],
        ["Superseded", incident["superseded_reason"] or "No"],
    ], [1.5, 5.1])
    doc.add_heading("Analyst assessment", level=1)
    doc.add_paragraph(clean(incident["analyst_disposition"] or "No separate assessment entered; workflow status is shown above."))
    doc.add_paragraph(clean(incident["notes"] or "No notes recorded"))
    doc.add_heading("Evidence timeline", level=1)
    _table(doc, ["Time UTC", "Source and actor", "Action and event"], [
        [event["timestamp"], f"{event['source']}\n{event['actor_id']}",
         f"{event['event_type']} / {event['action']}\n{event['event_id']}"] for event in events
    ], [1.55, 2.15, 2.9])
    missing = set(incident["evidence_ids"]) - {e["event_id"] for e in events}
    if missing:
        doc.add_paragraph("Unavailable evidence records: " + clean(", ".join(sorted(missing))))
    doc.add_heading("Evidence integrity", level=1)
    doc.add_paragraph("SHA-256 hashes identify the normalized evidence snapshot, including its stored raw reference. "
                      "They detect changes relative to this export; they are not digital signatures.")
    _table(doc, ["Event", "SHA 256"], [
        [e["event_id"], hashlib.sha256(json.dumps(e, sort_keys=True, separators=(",", ":"),
                                               ensure_ascii=True).encode()).hexdigest()]
        for e in events], [1.5, 5.1])
    return _bytes(doc)
