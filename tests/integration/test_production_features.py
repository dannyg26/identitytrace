import io
import json
import time
from concurrent.futures import ThreadPoolExecutor
from types import SimpleNamespace
from unittest.mock import Mock

import jwt
import pytest
from cryptography.hazmat.primitives.asymmetric import rsa
from docx import Document

from app.security import AccessPolicy
from tests.integration.test_ingestion_reliability import _chain


@pytest.fixture()
def oidc(client, monkeypatch):
    client.base_url = "https://testserver"
    for key, value in {
        "IDENTITYTRACE_ENV": "production", "IDENTITYTRACE_DEMO_MODE": "0",
        "IDENTITYTRACE_OIDC_ISSUER": "https://identity.example.test/tenant",
        "IDENTITYTRACE_OIDC_AUDIENCE": "identitytrace-api",
        "IDENTITYTRACE_OIDC_JWKS_URL": "https://identity.example.test/keys",
        "IDENTITYTRACE_ORGANIZATION_ID": "org-one",
        "IDENTITYTRACE_OIDC_ROLE_MAP": json.dumps({"IT.Viewer": "viewer", "IT.Admin": "admin"}),
    }.items():
        monkeypatch.setenv(key, value)
    policy = AccessPolicy()
    key = rsa.generate_private_key(public_exponent=65537, key_size=2048)
    policy.oidc.keys.get_signing_key_from_jwt = Mock(return_value=SimpleNamespace(key=key.public_key()))
    client.app.state.access_policy = policy

    def token(**overrides):
        claims = {"iss": policy.oidc.issuer, "aud": policy.oidc.audience, "sub": "person-one",
                  "iat": int(time.time()), "exp": int(time.time()) + 300,
                  "tid": "org-one", "roles": ["IT.Viewer"], **overrides}
        return jwt.encode(claims, key, algorithm="RS256", headers={"kid": "test-key"})

    return client, token, policy


def test_valid_external_token_is_role_limited(oidc):
    client, token, _ = oidc
    client.headers["Authorization"] = "Bearer " + token()
    assert client.get("/api/events").status_code == 200
    assert client.post("/api/events", json=_chain()[0]).status_code == 403
    assert client.get("/metrics").status_code == 403
    client.headers["Authorization"] = "Bearer " + token(roles=["IT.Admin"])
    assert client.get("/metrics").status_code == 200
    assert client.post("/api/events", json=_chain()[0]).status_code == 201


def test_production_has_no_local_credential_bypass(oidc):
    import hashlib
    client, _, policy = oidc
    policy.users = [{"name": "local", "role": "admin", "token_hash": hashlib.sha256(b"local-token").hexdigest()}]
    assert client.get("/api/events", headers={"Authorization": "Bearer local-token"}).status_code == 401
    assert client.get("/api/events", auth=("local", "password")).status_code == 401


def test_export_capacity_is_bounded(client, monkeypatch):
    monkeypatch.setattr("app.api.reports._slot", Mock(acquire=Mock(return_value=False)))
    response = client.get("/api/reports/summary.docx")
    assert response.status_code == 429
    assert response.headers["retry-after"] == "5"


def test_schema_verification_rejects_an_unmigrated_database(client):
    from sqlalchemy import text

    from app.models.db import engine, verify_schema
    verify_schema()
    with engine.begin() as connection:
        connection.execute(text("DELETE FROM identitytrace_schema_versions"))
    with pytest.raises(RuntimeError, match="migrate_database"):
        verify_schema()


@pytest.mark.parametrize("claims", [
    {"tid": "org-two"}, {"aud": "other-app"}, {"iss": "https://evil.test"},
    {"exp": 1}, {"iat": int(time.time()) + 3600}, {"roles": ["unknown"]},
    {"roles": "IT.Admin"}, {"roles": ["IT.Admin", "IT.Viewer"]}, {"sub": ""},
])
def test_external_tokens_fail_closed(oidc, claims):
    client, token, _ = oidc
    assert client.get("/api/events", headers={"Authorization": "Bearer " + token(**claims)}).status_code == 401


def test_signature_and_algorithm_tampering_rejected(oidc):
    client, token, policy = oidc
    bad = jwt.encode({"sub": "attacker"}, "a" * 32, algorithm="HS256", headers={"kid": "test"})
    assert client.get("/", headers={"Authorization": "Bearer " + bad}).status_code == 401
    policy.oidc.keys.get_signing_key_from_jwt.assert_not_called()
    parts = token().split(".")
    parts[2] = "AAAA"
    assert client.get("/", headers={"Authorization": "Bearer " + ".".join(parts)}).status_code == 401


def test_missing_claim_and_unavailable_keys_rejected(oidc):
    client, token, policy = oidc
    assert client.get("/", headers={"Authorization": "Bearer " + token(exp=None)}).status_code == 401
    policy.oidc.keys.get_signing_key_from_jwt.side_effect = jwt.PyJWKClientError("unavailable")
    response = client.get("/", headers={"Authorization": "Bearer " + token()})
    assert response.status_code == 401
    assert "unavailable" not in response.text


def test_production_forbids_demo_and_missing_oidc(monkeypatch):
    monkeypatch.setenv("IDENTITYTRACE_ENV", "production")
    with pytest.raises(ValueError, match="Production requires"):
        AccessPolicy()


def test_readiness_checks_database_and_config(client, monkeypatch):
    assert client.get("/ready").status_code == 200
    monkeypatch.setattr("app.main.SessionLocal", Mock(side_effect=RuntimeError("secret DSN")))
    response = client.get("/ready")
    assert response.status_code == 503
    assert "secret" not in response.text
    assert client.get("/health").status_code == 200


def test_metrics_and_logs_do_not_include_raw_ids_or_query_strings(client, caplog):
    response = client.get("/api/events/private-user-id?token=secret-token")
    assert len(response.headers["x-request-id"]) == 32
    metrics = client.get("/metrics").text
    assert "private-user-id" not in metrics and "secret-token" not in metrics
    assert "/api/events/{event_id}" in metrics
    entries = [json.loads(record.message) for record in caplog.records if record.name == "identitytrace.requests"]
    assert entries
    assert all("private-user-id" not in json.dumps(entry) and "secret-token" not in json.dumps(entry) for entry in entries)


def seed(client):
    for event in _chain():
        assert client.post("/api/events", json=event).status_code == 201
    return client.get("/api/incidents").json()[0]


def test_search_filters_and_exports_share_same_records(client):
    incident = seed(client)
    assert len(client.get("/api/incidents?q=alice").json()) == 1
    assert client.get("/api/incidents?q=not-present").json() == []
    assert client.get("/api/incidents?q=%25").json() == []  # literal %, not SQL wildcard
    assert client.get("/incidents?q=alice").status_code == 200
    assert "Download Word report" in client.get("/incidents").text
    assert client.get("/api/incidents?since=2027-01-01").json() == []
    assert client.get("/api/incidents?since=2027-01-01&until=2026-01-01").status_code == 422
    csv = client.get("/api/reports/incidents.csv?q=alice")
    assert incident["incident_id"] in csv.text
    empty = client.get("/api/reports/incidents.csv?q=not-present")
    assert incident["incident_id"] not in empty.text
    summary = Document(io.BytesIO(client.get("/api/reports/summary.docx?q=alice").content))
    assert incident["incident_id"] in " ".join(cell.text for row in summary.tables[0].rows for cell in row.cells)


def test_incident_docx_has_evidence_and_export_audit(client):
    incident = seed(client)
    response = client.get(f"/api/reports/incidents/{incident['incident_id']}.docx")
    assert response.status_code == 200
    doc = Document(io.BytesIO(response.content))
    assert doc.paragraphs[0].text == "IdentityTrace Incident Report"
    assert len(doc.tables[1].rows) == 4
    assert len(doc.tables[2].cell(1, 1).text) == 64
    audit = client.get("/api/audit").json()[0]
    assert audit["action"] == "report.exported"
    assert audit["after"]["incident_id"] == incident["incident_id"]
    assert client.get("/api/reports/incidents/missing.docx").status_code == 404


@pytest.mark.parametrize("assessment", ["", "Synthetic test completed; no real compromise."])
def test_resolved_assessment_is_consistent_in_page_and_report(client, assessment):
    incident = seed(client)
    incident_id = incident["incident_id"]
    response = client.patch(f"/api/incidents/{incident_id}", json={
        "status": "resolved", "analyst_disposition": assessment,
        "notes": "Synthetic acceptance exercise",
        "expected_updated_at": incident["updated_at"],
    })
    assert response.status_code == 200
    assert client.get(f"/api/incidents/{incident_id}").json()["analyst_disposition"] == assessment
    page = client.get(f"/incidents/{incident_id}").text
    assert 'id="disposition-input"' in page
    assert assessment in page
    doc = Document(io.BytesIO(client.get(f"/api/reports/incidents/{incident_id}.docx").content))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "No disposition recorded" not in text
    assert (assessment or "No separate assessment entered; workflow status is shown above.") in text
    assert "resolved" in " ".join(c.text for row in doc.tables[0].rows for c in row.cells)


def test_stale_analyst_edit_returns_conflict_without_losing_notes(client):
    incident = seed(client)
    url = f"/api/incidents/{incident['incident_id']}"
    first = client.patch(url, json={"notes": "First analyst", "expected_updated_at": incident["updated_at"]})
    assert first.status_code == 200
    second = client.patch(url, json={"notes": "Stale analyst", "expected_updated_at": incident["updated_at"]})
    assert second.status_code == 409
    assert client.get(url).json()["notes"] == "First analyst"


def test_csv_formula_injection_is_neutralized(client):
    incident = seed(client)
    from app.models.db import SessionLocal
    from app.models.incident import IncidentRecord
    with SessionLocal() as db:
        db.get(IncidentRecord, incident["incident_id"]).identity_id = "=HYPERLINK(unsafe)"
        db.commit()
    assert "'=HYPERLINK(unsafe)" in client.get("/api/reports/incidents.csv").text


def test_concurrent_duplicate_and_conflicting_ingestion_is_atomic(client):
    original = _chain()[0]
    conflicting = {**original, "actor_id": "different@example.test"}
    with ThreadPoolExecutor(max_workers=4) as pool:
        responses = list(pool.map(lambda event: client.post("/api/events", json=event),
                                  [original, conflicting] * 4))
    assert sorted(r.status_code for r in responses) == [201] * 4 + [409] * 4
    assert len(client.get("/api/events").json()) == 1
